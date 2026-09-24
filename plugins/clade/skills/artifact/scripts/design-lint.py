#!/usr/bin/env python3
"""design-lint — measure rendered design output against the Clade design floors.

The /frontend-design skill declares floors (WCAG 2.2 AA contrast, legibility
minimums, focal-surface caps, presentation word caps).  Most of them are
*rendered-output* rules: a clean grep over source proves nothing about them.
This is the validator that actually opens the artifact and measures.

Four lanes, each usable on its own:

  deck    <file.pptx>        source metrics via python-pptx (optional dep)
  render  <image|dir>        pixel metrics on rendered slides/pages (Pillow)
  html    <file.html|dir>    static analysis of artifact pages (stdlib only)
  source  <file|dir>         static analysis of a UI SOURCE tree — CSS/SCSS,
                             HTML, JSX/TSX, Vue, Svelte, Astro — against the
                             checkable half of the skill's design-rules.md:
                             spacing grid, type-scale caps, motion table,
                             token discipline, copy; plus .js/.ts for the
                             motion-runtime checks only (scroll-jacking,
                             uncapped DPR, a render loop with no visibility
                             gate) from signature-motion.md (stdlib only)

Exit status is 0 when every check passes, 1 when any FAIL is recorded, and 2
on a usage/environment error.  `--json` emits the full findings for machines.

Honesty contract: every check states what it can and cannot see.  Static HTML
analysis catches definite violations, never proves absence of all of them; the
render lane's contrast probe is a heuristic and is labelled as one.  A check
that cannot run reports SKIP with the reason — it never silently passes.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field, asdict
from html.parser import HTMLParser
from pathlib import Path
from statistics import median
from typing import Any, Iterable

# ─── Defaults (the floors declared in configs/skills/frontend-design) ───

BODY_MIN_PT = 13.0          # hard floor for essential presentation text
BODY_MEDIAN_PT = 18.0       # target median for audience-facing body copy
WORDS_PER_SLIDE = 70        # presentation word cap
FOCAL_COVERAGE_MAX = 0.45   # heavy full-bleed surface share of canvas
CONTRAST_TEXT = 4.5         # WCAG 2.2 AA, normal text
CONTRAST_LARGE = 3.0        # WCAG 2.2 AA, large text / non-text
BODY_MIN_PX = 16.0          # minimum body font-size on the web
INK_DENSITY_MAX = 0.60      # non-background coverage before a slide reads as a wall
TEXT_THINNESS_MIN = 0.30    # boundary/area ratio above which a colour reads as glyphs
TEXT_THINNESS_MAX = 0.98    # at/above this every pixel is an edge — anti-alias fringe
MIN_CANVAS_PX = 400         # smaller images are icons/logos, not design canvases

SEVERITY_ORDER = {"FAIL": 0, "WARN": 1, "SKIP": 2, "PASS": 3}


# ─── Findings ───


@dataclass
class Finding:
    check: str
    severity: str          # FAIL | WARN | SKIP | PASS
    target: str
    detail: str
    measured: Any = None
    threshold: Any = None


@dataclass
class Report:
    lane: str
    findings: list[Finding] = field(default_factory=list)

    def add(self, *a, **kw) -> None:
        self.findings.append(Finding(*a, **kw))

    @property
    def failed(self) -> bool:
        return any(f.severity == "FAIL" for f in self.findings)

    def render(self) -> str:
        rows = sorted(self.findings, key=lambda f: SEVERITY_ORDER.get(f.severity, 9))
        out = []
        for f in rows:
            mark = {"FAIL": "✗", "WARN": "!", "SKIP": "–", "PASS": "✓"}.get(f.severity, "?")
            out.append(f"  {mark} [{f.severity:4}] {f.target}: {f.detail}")
        counts: dict[str, int] = {}
        for f in self.findings:
            counts[f.severity] = counts.get(f.severity, 0) + 1
        summary = " ".join(f"{k}={v}" for k, v in sorted(counts.items()))
        out.append(f"  ── {self.lane}: {summary or 'no checks ran'}")
        return "\n".join(out)


# ─── Colour maths (WCAG 2.1/2.2 relative luminance) ───


def _channel(v: int) -> float:
    c = v / 255.0
    return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4


def luminance(rgb: tuple[int, int, int]) -> float:
    r, g, b = rgb
    return 0.2126 * _channel(r) + 0.7152 * _channel(g) + 0.0722 * _channel(b)


def contrast_ratio(a: tuple[int, int, int], b: tuple[int, int, int]) -> float:
    la, lb = luminance(a), luminance(b)
    if la < lb:
        la, lb = lb, la
    return (la + 0.05) / (lb + 0.05)


_HEX = re.compile(r"#([0-9a-fA-F]{3}|[0-9a-fA-F]{6})\b")
_RGB = re.compile(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)")
# Functions whose output is not the channels written inside them, and 8-digit
# hex, which carries alpha. Both must refuse to parse rather than be scraped.
_COMPUTED_COLOR = re.compile(
    r"\b(color-mix|light-dark|oklch|oklab|lab|lch|hwb|hsla?|color)\s*\(|#[0-9a-f]{8}\b", re.I)
_ALPHA = re.compile(r"rgba?\([^)]*[,/]\s*([\d.]+)\s*\)", re.I)

_NAMED = {
    "black": (0, 0, 0), "white": (255, 255, 255), "red": (255, 0, 0),
    "green": (0, 128, 0), "blue": (0, 0, 255), "gray": (128, 128, 128),
    "grey": (128, 128, 128), "silver": (192, 192, 192), "navy": (0, 0, 128),
}


_VAR_DEF = re.compile(r"(--[\w-]+)\s*:\s*([^;}]+)")
_VAR_USE = re.compile(r"var\(\s*(--[\w-]+)\s*(?:,\s*([^)]+))?\)")


def custom_properties(css: str) -> dict[str, str]:
    """Collect `--token: value` declarations so var() references can resolve.

    The skill tells designers to build on CSS custom properties, so a checker
    that gives up at `var(--ink)` is blind on exactly the idiom it recommends.
    Later definitions win, matching the cascade for same-specificity rules.
    """
    return {name: value.strip() for name, value in _VAR_DEF.findall(css)}


def resolve_value(value: str, props: dict[str, str], depth: int = 0) -> str:
    """Substitute var() references, honouring fallbacks, with a recursion cap."""
    if depth > 8 or "var(" not in value:
        return value

    def swap(m: re.Match) -> str:
        return props.get(m.group(1), (m.group(2) or "")).strip()

    return resolve_value(_VAR_USE.sub(swap, value), props, depth + 1)


def parse_color(text: str) -> tuple[int, int, int] | None:
    """Best-effort CSS colour → RGB. Returns None when not confidently parsed.

    Returning None is the *safe* answer, and the callers are built for it: an
    unresolved colour becomes an explicit SKIP, or the backdrop walk continues
    to the next opaque ancestor. Guessing is the dangerous answer — a bare
    `_HEX.search` on `color-mix(in srgb,#fff 12%,transparent)` finds `#fff` and
    reports a 12%-opacity veil as a solid white plate, inverting the verdict.
    """
    text = text.strip().lower()
    if text in _NAMED:
        return _NAMED[text]
    # Colour functions whose result is not the literal channels written inside
    # them. Never scrape a component colour out of one.
    if _COMPUTED_COLOR.search(text):
        return None
    m = _ALPHA.search(text)
    if m and float(m.group(1)) < 1.0:
        return None                                # translucent: needs compositing
    m = _HEX.search(text)
    if m:
        h = m.group(1)
        if len(h) == 3:
            h = "".join(c * 2 for c in h)
        return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]
    m = _RGB.search(text)
    if m:
        return (int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


# ─── Lane: deck (pptx source metrics) ───


def lint_deck(path: Path, report: Report, label: str) -> None:
    try:
        from pptx import Presentation           # type: ignore
        from pptx.util import Inches            # type: ignore
    except ImportError:
        report.add("deck.deps", "SKIP", label,
                   "python-pptx not installed — source metrics unavailable "
                   "(pip install python-pptx); render lane still applies")
        return

    prs = Presentation(str(path))
    canvas_w, canvas_h = prs.slide_width, prs.slide_height

    for idx, slide in enumerate(prs.slides, start=1):
        tag = f"{label} slide {idx}"
        texts: list[str] = []
        body_pt: list[float] = []

        for shape in slide.shapes:
            try:
                left, top = shape.left, shape.top
                width, height = shape.width, shape.height
            except (AttributeError, TypeError):
                left = top = width = height = None
            if None not in (left, top, width, height):
                if left < 0 or top < 0 or left + width > canvas_w or top + height > canvas_h:
                    report.add("deck.bounds", "FAIL", tag,
                               f"shape {shape.shape_id} extends outside the canvas")
            if not getattr(shape, "has_text_frame", False):
                continue
            texts.append(shape.text)
            in_body_zone = top is not None and Inches(2.0) <= top < Inches(6.0)
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text.strip() and run.font.size and in_body_zone:
                        body_pt.append(run.font.size.pt)

        words = re.findall(r"[A-Za-z0-9$%<+.'’-]+", " ".join(texts))
        if len(words) > WORDS_PER_SLIDE:
            report.add("deck.words", "FAIL", tag,
                       f"{len(words)} words exceeds the presentation cap",
                       measured=len(words), threshold=WORDS_PER_SLIDE)
        else:
            report.add("deck.words", "PASS", tag, f"{len(words)} words",
                       measured=len(words), threshold=WORDS_PER_SLIDE)

        if not body_pt:
            report.add("deck.type", "SKIP", tag,
                       "no explicitly-sized body-zone runs (inherited theme sizes are "
                       "not readable from the file — check the render lane)")
            continue
        lo, mid = min(body_pt), median(body_pt)
        if lo < BODY_MIN_PT:
            report.add("deck.type.floor", "FAIL", tag,
                       f"body-zone text at {lo:.1f}pt is below the legibility floor",
                       measured=round(lo, 1), threshold=BODY_MIN_PT)
        if mid < BODY_MEDIAN_PT:
            report.add("deck.type.median", "WARN", tag,
                       f"body-zone median {mid:.1f}pt is under the presentation target",
                       measured=round(mid, 1), threshold=BODY_MEDIAN_PT)
        if lo >= BODY_MIN_PT and mid >= BODY_MEDIAN_PT:
            report.add("deck.type", "PASS", tag,
                       f"min {lo:.1f}pt / median {mid:.1f}pt")


# ─── Lane: render (pixel metrics on rendered slides/pages) ───


def _pixels(image) -> Iterable:
    """Flat pixel sequence, across the Pillow API rename.

    Pillow 12 introduced `get_flattened_data()` and deprecated `getdata()`
    (removal in Pillow 14).  Deployments run whatever their distro ships, so
    probe rather than assume — picking one name strands the other environment
    with an AttributeError at the exact moment a check should be reporting.
    """
    getter = getattr(image, "get_flattened_data", None) or image.getdata
    return getter()


def _thinness(image, rgb: tuple[int, int, int]) -> float:
    """Fraction of this colour's pixels that lie on its own boundary.

    Glyph strokes are a few pixels wide, so nearly every pixel touches a
    non-matching neighbour and thinness approaches 1.  A filled panel is mostly
    interior and scores near 0.  This is what separates "text that must clear a
    contrast floor" from "a tinted surface that legitimately sits close to the
    page colour".
    """
    from PIL import Image, ImageFilter          # type: ignore

    # The image is already quantised, so exact equality is a safe mask test.
    bits = bytes(255 if px == rgb else 0 for px in _pixels(image))
    mask = Image.frombytes("L", image.size, bits)
    area = sum(1 for v in _pixels(mask) if v)
    if area == 0:
        return 0.0
    interior = mask.filter(ImageFilter.MinFilter(3))
    inner = sum(1 for v in _pixels(interior) if v)
    return 1.0 - (inner / area)


def _local_backdrop(image, rgb: tuple[int, int, int]):
    """The colour this one actually sits on, not the page's dominant colour.

    Text is judged against the surface directly behind it.  A caption on a
    yellow panel over a magenta page must clear its ratio against the yellow;
    comparing it to the page colour scores a slide that is in fact unreadable.
    Returns None when the ring around the colour is not readable.
    """
    from PIL import Image, ImageFilter          # type: ignore

    bits = bytes(255 if px == rgb else 0 for px in _pixels(image))
    mask = Image.frombytes("L", image.size, bits)
    ring = Image.frombytes(
        "L", image.size,
        bytes(255 if (d and not m) else 0
              for d, m in zip(_pixels(mask.filter(ImageFilter.MaxFilter(3))),
                              _pixels(mask))))
    counts: dict[tuple[int, int, int], int] = {}
    for px, on in zip(_pixels(image), _pixels(ring)):
        if on:
            counts[px] = counts.get(px, 0) + 1
    counts.pop(rgb, None)
    if not counts:
        return None
    return max(counts.items(), key=lambda kv: kv[1])[0]


def _dominant_colors(image, limit: int = 12) -> list[tuple[tuple[int, int, int], float]]:
    total = image.width * image.height
    counts = image.getcolors(maxcolors=total) or []
    counts.sort(key=lambda kv: kv[0], reverse=True)
    return [(rgb, n / total) for n, rgb in counts[:limit]]


def lint_render(path: Path, report: Report, label: str) -> None:
    try:
        from PIL import Image                   # type: ignore
    except ImportError:
        report.add("render.deps", "SKIP", label,
                   "Pillow not installed — no pixel metrics (pip install pillow)")
        return

    with Image.open(path) as probe:
        if probe.width < MIN_CANVAS_PX or probe.height < MIN_CANVAS_PX:
            report.add("render.scope", "SKIP", label,
                       f"{probe.width}×{probe.height} is below the canvas threshold — "
                       f"icons and logos are not judged as slides or pages")
            return

    with Image.open(path) as raw:
        image = raw.convert("RGB")
        # Quantise so anti-aliasing does not shatter the palette into thousands
        # of near-identical shades; keeps dominant-colour analysis meaningful.
        small = image.resize((min(image.width, 900), min(image.height, 900)))
        quant = small.quantize(colors=16, method=2).convert("RGB")

        total = image.width * image.height
        dark = 0
        for count, rgb in (image.getcolors(maxcolors=total) or []):
            # getcolors() yields a plain int per colour on greyscale ("L") and
            # palette images, not an (r, g, b) tuple — passing that straight to
            # luminance() indexes an int. Normalise instead of assuming RGB.
            if isinstance(rgb, (int, float)):
                rgb = (int(rgb), int(rgb), int(rgb))
            elif len(rgb) < 3:
                continue
            if luminance((rgb[0], rgb[1], rgb[2])) < 0.02:
                dark += count
        focal = dark / total

    if focal > FOCAL_COVERAGE_MAX:
        report.add("render.focal", "FAIL", label,
                   f"heavy surface covers {focal:.1%} of the canvas — it has stopped "
                   f"being emphasis and become the background",
                   measured=round(focal, 4), threshold=FOCAL_COVERAGE_MAX)
    else:
        report.add("render.focal", "PASS", label, f"heavy surface {focal:.1%}",
                   measured=round(focal, 4), threshold=FOCAL_COVERAGE_MAX)

    palette = _dominant_colors(quant)
    if not palette:
        report.add("render.contrast", "SKIP", label, "palette not readable")
        return
    background = palette[0][0]
    # Share alone cannot tell text from a tinted surface panel: a muted tile and a
    # paragraph can occupy the same fraction of the canvas, but only one of them
    # must clear a text-contrast floor.  Thinness separates them — glyph pixels sit
    # almost entirely on an edge, a solid panel is nearly all interior.
    detail = []
    for rgb, share in palette[1:]:
        if not 0.0005 <= share <= 0.15:
            continue
        thin = _thinness(quant, rgb)
        # thinness == 1.0 means the colour has no interior pixel at all: a 1px
        # scatter, which is what a renderer's anti-aliasing halo looks like.
        # Real glyph stems are 2px+ and always retain some interior.
        if TEXT_THINNESS_MIN <= thin < TEXT_THINNESS_MAX:
            detail.append((rgb, share, thin))
    if not detail:
        report.add("render.contrast", "SKIP", label,
                   "no text-like colour found (every minority colour reads as a solid "
                   "surface, not glyphs) — contrast unverified on pixels")
        return
    scored = []
    for rgb, share, thin in detail:
        backdrop = _local_backdrop(quant, rgb) or background
        scored.append((contrast_ratio(rgb, backdrop), rgb, share, thin, backdrop))
    worst, worst_rgb, worst_share, worst_thin, worst_bg = min(scored, key=lambda t: t[0])
    label = (f"dimmest text-like colour rgb{worst_rgb} ({worst_share:.2%} of canvas, "
             f"thinness {worst_thin:.2f}) on the surface behind it rgb{worst_bg} "
             f"= {worst:.2f}:1")
    # Deliberately capped at WARN.  Anti-aliasing fringe is thin and low-contrast by
    # construction, and no pixel-only test separates it from genuinely dim text.
    # A gate that cries wolf gets switched off, so the authoritative contrast lanes
    # are `html` (declared pairs) and the design system's own tokens; this probe
    # exists to point a human at the slide worth opening.
    if worst < CONTRAST_TEXT:
        report.add("render.contrast", "WARN", label,
                   f"{label} — under {CONTRAST_TEXT}:1. Heuristic: may be anti-aliasing "
                   f"fringe rather than text; confirm before treating as a defect",
                   measured=round(worst, 2), threshold=CONTRAST_TEXT)
    else:
        report.add("render.contrast", "PASS", label, label,
                   measured=round(worst, 2), threshold=CONTRAST_TEXT)


# ─── Lane: html (static analysis of artifact pages) ───


_STYLE_BLOCK = re.compile(r"<style[^>]*>(.*?)</style>", re.S | re.I)
_IMG_TAG = re.compile(r"<img\b[^>]*>", re.I)
_HEADING = re.compile(r"<h([1-6])\b", re.I)
_DECL = re.compile(r"([a-z-]+)\s*:\s*([^;{}]+)", re.I)
_RULE = re.compile(r"([^{}]+)\{([^{}]*)\}", re.S)
_FONT_SIZE_PX = re.compile(r"font-size\s*:\s*([\d.]+)px", re.I)

# ─── Inherited-contrast cascade (the inversion-band blind spot) ───
#
# Pairing colour+background *within one rule* misses the most common real
# failure: an inversion band (`.thesis{background:var(--fg);color:var(--bg)}`)
# containing a child that repaints only its own background
# (`blockquote{background:var(--muted)}`). The child's text colour is
# inherited from the band, its background comes from the un-inverted palette,
# and the two collide — measured at 1.17:1 on a real published artifact while
# the declared-pair check reported no failure. Resolving that needs the HTML
# tree, not the stylesheet alone.

_UNSUPPORTED_SEL = re.compile(r"[>+~*\[]|::|:(?!root\b)")
_COMPOUND = re.compile(r"^([a-z][\w-]*)?((?:[.#][\w-]+)*)$", re.I)
# Elements that never render inherited text of their own.
_VOID_OR_META = {"style", "script", "head", "meta", "link", "title", "br", "img",
                 "hr", "input", "source", "path", "svg", "circle", "rect"}


def _split_media(css: str) -> tuple[str, str]:
    """Split CSS into (base, prefers-color-scheme:dark) sources.

    Without this, `custom_properties`'s last-definition-wins collapses a
    two-theme page onto whichever palette is declared last — so one whole
    theme is never measured.
    """
    base: list[str] = []
    dark: list[str] = []
    i = 0
    while True:
        m = re.compile(r"@media[^{]*prefers-color-scheme\s*:\s*dark[^{]*\{", re.I).search(css, i)
        if not m:
            base.append(css[i:])
            return "".join(base), "".join(dark)
        base.append(css[i:m.start()])
        depth, j = 1, m.end()
        while j < len(css) and depth:
            if css[j] == "{":
                depth += 1
            elif css[j] == "}":
                depth -= 1
            j += 1
        dark.append(css[m.end():j - 1])
        i = j


class _Tree(HTMLParser):
    """Minimal element tree: (tag, id, classes, parent, has_text)."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.nodes: list[dict] = []
        self._stack: list[int] = []

    def handle_starttag(self, tag: str, attrs: list) -> None:
        a = dict(attrs)
        self.nodes.append({
            "tag": tag,
            "id": a.get("id") or "",
            "classes": set((a.get("class") or "").split()),
            "parent": self._stack[-1] if self._stack else None,
            "text": False,
        })
        if tag not in _VOID_OR_META:
            self._stack.append(len(self.nodes) - 1)

    def handle_endtag(self, tag: str) -> None:
        for k in range(len(self._stack) - 1, -1, -1):
            if self.nodes[self._stack[k]]["tag"] == tag:
                del self._stack[k:]
                return

    def handle_data(self, data: str) -> None:
        if data.strip() and self._stack:
            node = self.nodes[self._stack[-1]]
            if node["tag"] not in _VOID_OR_META:
                node["text"] = True


def _match_compound(node: dict, comp: str) -> bool:
    m = _COMPOUND.match(comp)
    if not m:
        return False
    tag, rest = m.group(1), m.group(2) or ""
    if tag and tag.lower() != node["tag"]:
        return False
    for token in re.findall(r"[.#][\w-]+", rest):
        if token[0] == "." and token[1:] not in node["classes"]:
            return False
        if token[0] == "#" and token[1:] != node["id"]:
            return False
    return True


def _matches(nodes: list[dict], idx: int, selector: str) -> bool:
    """Descendant-combinator matching, right to left."""
    parts = selector.split()
    if not parts or not _match_compound(nodes[idx], parts[-1]):
        return False
    cur = nodes[idx]["parent"]
    for comp in reversed(parts[:-1]):
        while cur is not None and not _match_compound(nodes[cur], comp):
            cur = nodes[cur]["parent"]
        if cur is None:
            return False
        cur = nodes[cur]["parent"]
    return True


def _specificity(selector: str) -> tuple[int, int, int]:
    return (selector.count("#"), selector.count("."),
            len(re.findall(r"(?:^|\s)[a-z][\w-]*", selector, re.I)))


def inherited_contrast(text: str, css: str, tag: str, report: Report) -> None:
    """Measure each text node's inherited colour against its effective backdrop."""
    tree = _Tree()
    try:
        tree.feed(text)
    except Exception as exc:                                   # malformed markup
        report.add("html.contrast.inherited", "SKIP", tag, f"unparseable markup: {exc}")
        return
    nodes = tree.nodes
    if not nodes:
        report.add("html.contrast.inherited", "SKIP", tag, "no elements parsed")
        return

    base_css, dark_css = _split_media(css)
    themes = [("light", custom_properties(base_css))]
    if dark_css.strip():
        dark_props = dict(themes[0][1])
        dark_props.update(custom_properties(dark_css))
        themes.append(("dark", dark_props))

    # Rules in cascade order, skipping selectors this matcher cannot honour.
    rules, skipped = [], 0
    for order, (selector, body) in enumerate(_RULE.findall(base_css)):
        decls = {k.lower().strip(): v.strip() for k, v in _DECL.findall(body)}
        relevant = {"color", "background", "background-color"} & decls.keys()
        if not relevant and not any(k.startswith("--") for k in decls):
            continue
        for sel in (s.strip() for s in selector.split(",")):
            if not sel or sel.startswith("@") or _UNSUPPORTED_SEL.search(sel):
                skipped += 1
                continue
            rules.append((_specificity(sel), order, sel, decls))
    rules.sort(key=lambda r: (r[0], r[1]))

    matched: list[dict] = [{} for _ in nodes]
    for _, _, sel, decls in rules:
        for i in range(len(nodes)):
            if _matches(nodes, i, sel):
                matched[i].update(decls)

    def scoped_props(i: int, base: dict) -> dict:
        """Custom properties inherit, so a band may re-point --muted for its
        own subtree. Resolving every var() against the :root palette would
        report the *unfixed* colour on a page that scopes its tokens correctly.
        """
        chain = []
        j: int | None = i
        while j is not None:
            chain.append(j)
            j = nodes[j]["parent"]
        props = dict(base)
        for k in reversed(chain):                       # root → self, nearest wins
            props.update({n: v for n, v in matched[k].items() if n.startswith("--")})
        return props

    def inherited_color(i: int | None, base: dict) -> tuple[int, int, int] | None:
        while i is not None:
            raw = matched[i].get("color")
            if raw:
                # A var() in an inherited property resolves in the scope of the
                # element that declared it, not the one that inherits it.
                c = parse_color(resolve_value(raw, scoped_props(i, base)))
                if c:
                    return c
            i = nodes[i]["parent"]
        return None

    def effective_bg(i: int | None, base: dict) -> tuple[int, int, int] | None:
        while i is not None:
            raw = matched[i].get("background-color") or matched[i].get("background")
            if raw and not re.search(r"\b(transparent|none)\b", raw, re.I):
                c = parse_color(resolve_value(raw, scoped_props(i, base)))
                if c:
                    return c
            i = nodes[i]["parent"]
        return None

    worst: tuple[float, str] | None = None
    for theme, props in themes:
        for i, node in enumerate(nodes):
            if not node["text"]:
                continue
            fg, bg = inherited_color(i, props), effective_bg(i, props)
            if not fg or not bg:
                continue
            ratio = contrast_ratio(fg, bg)
            if worst is None or ratio < worst[0]:
                where = node["tag"] + ("." + sorted(node["classes"])[0] if node["classes"] else "")
                worst = (ratio, f"<{where}> rgb{fg} on rgb{bg} = {ratio:.2f}:1 [{theme}]")

    if worst is None:
        report.add("html.contrast.inherited", "SKIP", tag,
                   "no text node resolved to both an inherited colour and an opaque backdrop")
        return
    note = f" ({skipped} unsupported selector(s) not matched)" if skipped else ""
    if worst[0] < CONTRAST_LARGE:
        report.add("html.contrast.inherited", "FAIL", tag,
                   f"worst inherited pair {worst[1]} — below {CONTRAST_LARGE}:1 even for "
                   f"large text{note}", measured=round(worst[0], 2), threshold=CONTRAST_TEXT)
    elif worst[0] < CONTRAST_TEXT:
        report.add("html.contrast.inherited", "WARN", tag,
                   f"worst inherited pair {worst[1]} — clears large-text {CONTRAST_LARGE}:1 "
                   f"but not body {CONTRAST_TEXT}:1{note}",
                   measured=round(worst[0], 2), threshold=CONTRAST_TEXT)
    else:
        report.add("html.contrast.inherited", "PASS", tag,
                   f"worst inherited pair {worst[1]}{note}",
                   measured=round(worst[0], 2), threshold=CONTRAST_TEXT)


def lint_html(path: Path, report: Report, label: str) -> None:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        report.add("html.read", "SKIP", label, f"unreadable: {exc}")
        return

    tag = label
    css = "\n".join(_STYLE_BLOCK.findall(text))
    inline_styles = re.findall(r'style\s*=\s*"([^"]*)"', text, re.I)
    all_css = css + "\n" + "\n".join(inline_styles)

    # -- motion respects preference ------------------------------------------
    has_motion = bool(re.search(r"(animation|transition)\s*:", all_css, re.I)) or \
        bool(re.search(r"@keyframes", css, re.I))
    guarded = bool(re.search(r"prefers-reduced-motion", text, re.I))
    if has_motion and not guarded:
        report.add("html.motion", "FAIL", tag,
                   "declares animation/transition with no prefers-reduced-motion guard")
    elif has_motion:
        report.add("html.motion", "PASS", tag, "motion is guarded by prefers-reduced-motion")

    # -- focus visibility -----------------------------------------------------
    kills_outline = re.search(r"outline\s*:\s*(none|0)\b", all_css, re.I)
    restores = _FOCUS_RESTORE.search(css)
    if kills_outline and not restores:
        report.add("html.focus", "FAIL", tag,
                   "removes the focus outline without providing a replacement indicator")
    elif kills_outline:
        report.add("html.focus", "PASS", tag, "outline removed but a focus style is restored")

    # -- images carry alt -----------------------------------------------------
    imgs = _IMG_TAG.findall(text)
    missing_alt = [t for t in imgs if not re.search(r"\balt\s*=", t, re.I)]
    if missing_alt:
        report.add("html.alt", "FAIL", tag,
                   f"{len(missing_alt)} of {len(imgs)} <img> tags have no alt attribute",
                   measured=len(missing_alt), threshold=0)
    elif imgs:
        report.add("html.alt", "PASS", tag, f"all {len(imgs)} <img> tags carry alt")

    # -- heading order --------------------------------------------------------
    levels = [int(m) for m in _HEADING.findall(text)]
    if levels:
        h1s = levels.count(1)
        problems = []
        if h1s == 0:
            problems.append("no <h1>")
        elif h1s > 1:
            problems.append(f"{h1s} <h1> elements")
        prev = None
        for lv in levels:
            if prev is not None and lv > prev + 1:
                problems.append(f"level jump h{prev}→h{lv}")
                break
            prev = lv
        if problems:
            report.add("html.headings", "WARN", tag,
                       "heading structure: " + ", ".join(problems))
        else:
            report.add("html.headings", "PASS", tag,
                       f"{len(levels)} headings, single h1, no skipped levels")

    # -- declared colour pairs ------------------------------------------------
    # A pair that cannot be statically resolved (unsupported colour function,
    # a custom property with no fallback) must never be silently dropped: if
    # every OTHER pair happens to pass, the unresolved one would otherwise
    # vanish into a bare PASS that never actually checked it.
    props = custom_properties(all_css)
    worst: tuple[float, str] | None = None
    unresolved: list[str] = []
    for selector, body in _RULE.findall(css):
        decls = {k.lower().strip(): v.strip() for k, v in _DECL.findall(body)}
        fg = decls.get("color")
        bg = decls.get("background-color") or decls.get("background")
        if not fg or not bg:
            continue
        sel = " ".join(selector.split())[:60]
        r_fg, r_bg = resolve_value(fg, props), resolve_value(bg, props)
        c_fg, c_bg = parse_color(r_fg), parse_color(r_bg)
        if not c_fg or not c_bg:
            unresolved.append(f"{sel} — {fg.strip()} on {bg.strip()}")
            continue
        ratio = contrast_ratio(c_fg, c_bg)
        if worst is None or ratio < worst[0]:
            worst = (ratio, f"{sel} — {r_fg.strip()} on {r_bg.strip()} = {ratio:.2f}:1")
    if worst is None:
        if unresolved:
            report.add("html.contrast", "SKIP", tag,
                       f"{len(unresolved)} rule(s) declare colour+background that could "
                       f"not be statically resolved (e.g. {unresolved[0]}) — unsupported "
                       f"colour function or a custom property with no fallback")
        else:
            report.add("html.contrast", "SKIP", tag,
                       "no rule declares colour and background together — static analysis "
                       "cannot resolve inherited or computed pairs")
    elif worst[0] < CONTRAST_TEXT:
        sev = "FAIL" if worst[0] < CONTRAST_LARGE else "WARN"
        report.add("html.contrast", sev, tag, worst[1] + " (AA text needs 4.5:1)",
                   measured=round(worst[0], 2), threshold=CONTRAST_TEXT)
    elif unresolved:
        report.add("html.contrast", "SKIP", tag,
                   f"worst resolvable pair clears {CONTRAST_TEXT}:1 ({worst[1]}), but "
                   f"{len(unresolved)} other pair(s) could not be statically resolved "
                   f"(e.g. {unresolved[0]}) and were never verified",
                   measured=round(worst[0], 2), threshold=CONTRAST_TEXT)
    else:
        report.add("html.contrast", "PASS", tag, "worst declared pair " + worst[1],
                   measured=round(worst[0], 2), threshold=CONTRAST_TEXT)

    # -- inherited colour vs effective backdrop (needs the DOM, not just CSS) --
    inherited_contrast(text, css, tag, report)

    # -- body type size -------------------------------------------------------
    sizes = [float(s) for s in _FONT_SIZE_PX.findall(all_css)]
    tiny = [s for s in sizes if s < BODY_MIN_PX]
    if tiny:
        report.add("html.type", "WARN", tag,
                   f"{len(tiny)} font-size declarations below {BODY_MIN_PX:.0f}px "
                   f"(smallest {min(tiny):.0f}px) — fine for metadata, not for body copy",
                   measured=min(tiny), threshold=BODY_MIN_PX)
    elif sizes:
        report.add("html.type", "PASS", tag,
                   f"{len(sizes)} font-size declarations, smallest {min(sizes):.0f}px")


# ─── Lane: source (static analysis of UI source trees) ───
#
# The html lane judges one rendered artifact page. This lane judges the SOURCE
# a product is built from — CSS/SCSS/LESS, HTML, and the template/JSX files
# that carry class names and copy — against the checkable half of
# configs/skills/frontend-design/references/design-rules.md: the spacing grid,
# the type-scale caps, the motion table, token discipline and the copy ban.
# Those rules exist because an agent asked for "premium, polished, designed"
# ships 13/18/22/27 px spacing, nine font sizes and a 900 ms fade by default;
# an explicit system it executes well.
#
# What it cannot see (one primary action per viewport, shared edges, tier
# contrast, control heights, responsive re-composition, every state) is
# reported as a SKIP naming the screenshot review, so a clean run never reads
# as a clean page.
#
# Severity policy: FAIL only for the two accessibility floors the html lane
# already fails on (motion with no reduced-motion guard, focus outline removed
# without a replacement). Everything else is WARN with the measured value,
# because each of those rules admits a written visual reason — the reviewer
# decides, this instrument names the sites. The thresholds are the essay's own
# numbers (4-px grid, 4–6 sizes, 3–4 weights, at most 2 families, 400 ms,
# 16 px), not measured fire rates; run it on a tree you consider clean before
# treating a WARN count as a verdict.

SOURCE_SUFFIXES = {".css", ".scss", ".less", ".html", ".htm",
                   ".tsx", ".jsx", ".vue", ".svelte", ".astro"}
SOURCE_MARKUP_SUFFIXES = {".html", ".htm", ".vue", ".svelte", ".astro"}
SOURCE_SCRIPT_SUFFIXES = {".tsx", ".jsx"}
# Plain script files carry the scroll story's runtime (GSAP, canvas, Three.js)
# and get ONLY the motion-runtime checks — never spacing, type or colour: a
# `padding: 13` inside a PDF library is not a design decision, and the guard
# that matters (`useReducedMotion`, `matchMedia('(prefers-reduced-motion')`)
# lives here as often as in a stylesheet.
SOURCE_RUNTIME_SUFFIXES = {".js", ".mjs", ".cjs", ".ts", ".mts", ".cts"}
SOURCE_SKIP_DIRS = {"node_modules", "dist", "build", "out", "coverage", "vendor",
                    ".git", ".next", ".nuxt", ".svelte-kit", ".turbo",
                    "storybook-static", "__pycache__"}
SPACING_GRID_PX = 4
SPACING_NUDGE_PX = 2            # 1–2 px is a hairline or an optical nudge, not space
TYPE_SIZES_MAX = 6
TYPE_WEIGHTS_MAX = 4
TYPE_FAMILIES_MAX = 2
RADIUS_VALUES_MAX = 4
SHADOW_VALUES_MAX = 3
MOTION_MAX_MS = 400.0           # above this only a signature sequence is allowed
MOTION_TRANSLATE_MAX_PX = 16.0
EXAMPLES_SHOWN = 5

SLOP_PHRASES = (
    "revolutioni[sz]e", "unlock the power", "transform your business",
    "get started today", "next[- ]generation", "supercharge", "seamlessly",
    "cutting[- ]edge", "game[- ]changing", "unleash", "赋能", "重新定义", "颠覆",
)
_SLOP = re.compile("(?<![A-Za-z])(?:" + "|".join(SLOP_PHRASES) + ")", re.I)

_GENERIC_FAMILIES = {"serif", "sans-serif", "monospace", "cursive", "fantasy",
                     "system-ui", "ui-sans-serif", "ui-serif", "ui-monospace",
                     "ui-rounded", "inherit", "initial", "unset", "revert", "math",
                     "emoji", "fangsong"}
_TW_FONT_SIZE_PX = {"xs": 12.0, "sm": 14.0, "base": 16.0, "lg": 18.0, "xl": 20.0,
                    "2xl": 24.0, "3xl": 30.0, "4xl": 36.0, "5xl": 48.0, "6xl": 60.0,
                    "7xl": 72.0, "8xl": 96.0, "9xl": 128.0}
_TW_FONT_WEIGHT = {"thin": 100, "extralight": 200, "light": 300, "normal": 400,
                   "medium": 500, "semibold": 600, "bold": 700, "extrabold": 800,
                   "black": 900}
_CSS_WEIGHT = {"normal": 400, "bold": 700}
_LAYOUT_TRANSITION_PROPS = {"all", "width", "height", "top", "left", "right",
                            "bottom", "margin", "padding", "inset", "min-width",
                            "max-width", "min-height", "max-height", "font-size",
                            "line-height"}
# An infinite animation is a loader when its keyframe name or its selector says
# so; anything else that loops forever is the "background drifting to no
# purpose" the rules forbid.
_LOADER_HINT = re.compile(
    r"spin|load|skeleton|shimmer|pulse|progress|ping|indeterminate|busy|wait", re.I)

_SRC_COMMENT = re.compile(r"/\*.*?\*/|<!--.*?-->", re.S)
_SRC_LINE_COMMENT = re.compile(r"^[ \t]*//[^\n]*$", re.M)
_SRC_SCRIPT_BLOCK = re.compile(r"<script\b[^>]*>.*?</script>", re.S | re.I)
_SRC_TAG = re.compile(r"<[^>]*>")
# `name: value` in CSS, SCSS, style attributes, CSS-in-JS template literals and
# React style objects (camelCase is normalised). Custom-property definitions
# (`--ink: #222`) never match: the name must start with a letter and cannot be
# preceded by `-`, so a token file is not reported as a wall of literals.
_PROP_DECL = re.compile(
    r"(?<![\w$-])([A-Za-z][\w-]*)\s*:\s*('[^'\n]*'|\"[^\"\n]*\"|[^;{}\n]+)")
_LENGTH = re.compile(r"(-?\d*\.?\d+)(px|rem)\b")
_TIME = re.compile(r"(\d*\.?\d+)(ms|s)\b")
_COLOR_LITERAL = re.compile(r"#[0-9a-fA-F]{3,8}\b|\b(?:rgb|hsl)a?\(", re.I)
_UNRESOLVABLE = re.compile(r"\b(?:calc|clamp|min|max|env|var)\(", re.I)
_KEYFRAMES_OPEN = re.compile(r"@keyframes\s+([\w-]+)\s*\{", re.I)
_TRANSLATE = re.compile(r"translate(?:X|Y|3d)?\(\s*(-?\d*\.?\d+)px", re.I)
# A focus rule counts as a RESTORE only when it sets an indicator. The obvious
# regex — any :focus rule mentioning `outline` — is satisfied by the very rule
# that removes it (`button:focus { outline: none }`), which is the one case
# the check exists to catch; the html lane shipped exactly that blind spot.
_FOCUS_RESTORE = re.compile(
    r":focus(?:-visible|-within)?\b[^{]*\{[^}]*"
    r"(?:outline(?!\s*:\s*(?:none|0)\b)|box-shadow|border|ring)", re.I | re.S)
_MOTION_GUARD = re.compile(r"prefers-reduced-motion|useReducedMotion|reduced-?motion", re.I)

_TW_SPACING = re.compile(
    r"(?<![\w-])-?(?:p|px|py|pt|pr|pb|pl|ps|pe|m|mx|my|mt|mr|mb|ml|ms|me|gap|gap-x|gap-y|"
    r"space-x|space-y)-\[(-?\d*\.?\d+)px\]")
_TW_TEXT_PX = re.compile(r"(?<![\w-])text-\[(\d*\.?\d+)px\]")
_TW_TEXT_STEP = re.compile(r"(?<![\w-])text-(xs|sm|base|lg|xl|[2-9]xl)(?![\w-])")
_TW_WEIGHT = re.compile(
    r"(?<![\w-])font-(thin|extralight|light|normal|medium|semibold|bold|extrabold|black)(?![\w-])")
_TW_RADIUS = re.compile(
    r"(?<![\w-])rounded(?:-(?:t|r|b|l|tl|tr|br|bl|s|e|ss|se|es|ee))?-\[(\d*\.?\d+)px\]")
_TW_DURATION = re.compile(r"(?<![\w-])duration-(?:\[(\d+)ms\]|(\d+))(?![\w-])")
_TW_TRANSITION_ALL = re.compile(r"(?<![\w-])transition-all(?![\w-])")
# Only Tailwind's own infinite utilities can be judged statically: spin, ping
# and pulse are loaders, bounce is ambient. A custom `animate-<name>` (an
# `animate-in` from tailwindcss-animate is a one-shot entrance) is motion, but
# whether it loops lives in the config this scan does not read.
_TW_ANIMATE = re.compile(r"(?<![\w-])animate-(spin|ping|pulse|bounce)(?![\w-])")
_TW_MOTION = re.compile(r"(?<![\w-])(?:transition(?:-[\w]+)?|duration-[\w\[\]]+|animate-[\w\[\]]+)(?![\w-])")
_TW_COLOR = re.compile(
    r"(?<![\w-])(?:bg|text|border|from|via|to|ring|fill|stroke|outline|decoration|divide|"
    r"shadow|accent|caret|placeholder)-\[#[0-9a-fA-F]{3,8}\]")
_TW_OUTLINE_NONE = re.compile(r"(?<![\w-])outline-none(?![\w-])")
_TW_FOCUS_RESTORE = re.compile(r"(?<![\w-])focus(?:-visible|-within)?:(?:ring|outline|border|shadow)")

# ─── Motion runtime (signature-motion.md §5) ───
#
# Three properties of a scroll story or hero piece that a static read CAN see.
# Scroll-jacking: a wheel/touch listener that calls preventDefault, or a
# page-snapping library. DPR: `devicePixelRatio` used with no Math.min/clamp
# nearby — the 2–3x canvas that drops frames on phones. Render loop: a
# requestAnimationFrame loop or setAnimationLoop with no visibility gate
# anywhere in the tree, so the canvas keeps drawing off-screen and after the
# intro. Each is a WARN naming the site; a legitimate drag handler or a demand
# frameloop is the reviewer's call, not this instrument's.
_WHEEL_LISTENER = re.compile(
    r"addEventListener\s*\(\s*['\"](?:wheel|mousewheel|DOMMouseScroll|touchmove)['\"]", re.I)
_PREVENT_DEFAULT = re.compile(r"\.preventDefault\s*\(")
_SCROLLJACK_LIB = re.compile(
    r"['\"](?:fullpage\.js|fullpage|@fullpage/[\w-]+|pagepiling(?:\.js)?|jquery\.scrollify|"
    r"scrollify|onepage-scroll)['\"]", re.I)
_DPR_USE = re.compile(r"devicePixelRatio")
_DPR_CAP = re.compile(r"Math\.min\s*\(|clamp\s*\(|dpr\s*[=:]\s*\{?\s*\[|setPixelRatio\s*\(\s*Math\.min", re.I)
_RENDER_LOOP = re.compile(r"\bsetAnimationLoop\s*\(|\brequestAnimationFrame\s*\(")
_VISIBILITY_GATE = re.compile(
    r"IntersectionObserver|visibilitychange|document\.hidden|isIntersecting|\bonEnter\b|"
    r"\bonLeave\b|\bonToggle\b|\binView\b|\buseInView\b|frameloop\s*=\s*['\"]demand", re.I)


@dataclass
class _Site:
    """One place in the tree where a rule was measured: `file:line what`."""
    label: str
    line: int
    text: str

    def __str__(self) -> str:
        return f"{self.label}:{self.line} {self.text}"


@dataclass
class _SourceScan:
    spacing_total: int = 0
    spacing_off: list[_Site] = field(default_factory=list)
    margins: int = 0
    gaps: int = 0
    sizes: dict[float, _Site] = field(default_factory=dict)
    weights: dict[int, _Site] = field(default_factory=dict)
    families: dict[str, _Site] = field(default_factory=dict)
    radii: dict[str, _Site] = field(default_factory=dict)
    shadows: dict[str, _Site] = field(default_factory=dict)
    durations: int = 0
    durations_over: list[_Site] = field(default_factory=list)
    layout_transitions: list[_Site] = field(default_factory=list)
    infinite_ambient: list[_Site] = field(default_factory=list)
    keyframes: int = 0
    translate_over: list[_Site] = field(default_factory=list)
    colors: list[_Site] = field(default_factory=list)
    slop: list[_Site] = field(default_factory=list)
    markup_files: int = 0
    has_motion: bool = False
    has_guard: bool = False
    kills_outline: list[_Site] = field(default_factory=list)
    restores_focus: bool = False
    script_files: int = 0
    scrolljack: list[_Site] = field(default_factory=list)
    dpr_uncapped: list[_Site] = field(default_factory=list)
    dpr_capped: int = 0
    render_loops: list[_Site] = field(default_factory=list)
    has_visibility_gate: bool = False


def _blank_comments(text: str) -> str:
    """Replace comment bodies with spaces, keeping every newline so line
    numbers computed on the result still point at the original file."""
    def blank(m: re.Match) -> str:
        return re.sub(r"[^\n]", " ", m.group(0))
    return _SRC_LINE_COMMENT.sub(blank, _SRC_COMMENT.sub(blank, text))


def _line_of(text: str, pos: int) -> int:
    return text.count("\n", 0, pos) + 1


def _kebab(name: str) -> str:
    return re.sub(r"(?<!^)(?=[A-Z])", "-", name).lower()


def _px(value: str, unit: str) -> float:
    return float(value) * (16.0 if unit == "rem" else 1.0)


def _off_grid(px: float) -> bool:
    if abs(px) <= SPACING_NUDGE_PX:
        return False
    return abs(px / SPACING_GRID_PX - round(px / SPACING_GRID_PX)) > 1e-6


def _first(sites: list[_Site]) -> str:
    shown = ", ".join(str(s) for s in sites[:EXAMPLES_SHOWN])
    more = len(sites) - EXAMPLES_SHOWN
    return shown + (f" … +{more} more" if more > 0 else "")


def _scan_declarations(text: str, label: str, props: dict[str, str], scan: _SourceScan) -> None:
    for m in _PROP_DECL.finditer(text):
        name = _kebab(m.group(1))
        raw = m.group(2).strip()
        # A React style object or CSS-in-JS value is a quoted string; the
        # declaration regex stops at the line end, so trim to the closing quote
        # or `color: '#fff'` on the same line is read as part of `paddingTop`.
        if raw[:1] in ("'", '"') and raw.find(raw[0], 1) > 0:
            raw = raw[1:raw.find(raw[0], 1)]
        if not raw or "$" in raw:              # template interpolation: unresolvable
            continue
        line = _line_of(text, m.start())
        site = _Site(label, line, f"{name}: {raw[:48]}")
        value = resolve_value(raw, props)

        if name in ("gap", "row-gap", "column-gap") or name.startswith(("margin", "padding")):
            if name.startswith("margin"):
                scan.margins += 1
            elif name.startswith("gap") or name.endswith("gap"):
                scan.gaps += 1
            if _UNRESOLVABLE.search(value):
                continue
            lengths = [_px(lm.group(1), lm.group(2)) for lm in _LENGTH.finditer(value)]
            if lengths:
                scan.spacing_total += 1
                if any(_off_grid(px) for px in lengths):
                    scan.spacing_off.append(site)
        elif name == "font-size":
            sm = _LENGTH.search(value)
            if sm and not _UNRESOLVABLE.search(value):
                scan.sizes.setdefault(round(_px(sm.group(1), sm.group(2)), 1), site)
        elif name == "font-weight":
            wm = re.search(r"\b([1-9]00)\b", value)
            weight = int(wm.group(1)) if wm else _CSS_WEIGHT.get(value.strip().lower())
            if weight:
                scan.weights.setdefault(weight, site)
        elif name == "font-family":
            first = value.split(",")[0].strip().strip("'\"").lower()
            if first and first not in _GENERIC_FAMILIES and "var(" not in first:
                scan.families.setdefault(first, site)
        elif name == "border-radius" or name.endswith("-radius"):
            key = "pill" if re.search(r"%|9{3,}|1e\d|vmax", value) else " ".join(value.split())
            if key and not _UNRESOLVABLE.search(key):
                scan.radii.setdefault(key, site)
        elif name == "box-shadow":
            key = " ".join(value.split())
            if key and key.lower() != "none":
                scan.shadows.setdefault(key, site)
        elif name in ("transition", "transition-property", "transition-duration",
                      "animation", "animation-duration", "animation-iteration-count"):
            _scan_motion_decl(text, m, name, value, site, scan)
        elif name in ("outline", "outline-style") and re.match(r"(none|0)\b", value.strip()):
            scan.kills_outline.append(site)

        # Shadow colours are judged under geometry.shadow; counting them here
        # too would report one shadow token as two findings.
        if _COLOR_LITERAL.search(raw) and name not in (
                "content", "src", "href", "url", "background-image", "mask-image",
                "box-shadow", "text-shadow", "filter"):
            scan.colors.append(site)


def _scan_motion_decl(text: str, m: re.Match, name: str, value: str,
                      site: _Site, scan: _SourceScan) -> None:
    if value.strip().lower() in ("none", "0", "0s", "0ms"):
        return
    scan.has_motion = True
    context = text[max(0, m.start() - 120):m.end()]
    loader = bool(_LOADER_HINT.search(context))
    if name in ("animation", "animation-iteration-count") and "infinite" in value.lower() and not loader:
        scan.infinite_ambient.append(site)
    for segment in value.split(","):
        tm = _TIME.search(segment)
        if tm and name != "animation-iteration-count":
            scan.durations += 1
            ms = float(tm.group(1)) * (1000.0 if tm.group(2) == "s" else 1.0)
            if ms > MOTION_MAX_MS and not (loader and "infinite" in value.lower()):
                scan.durations_over.append(site)
        if name in ("transition", "transition-property"):
            head = re.match(r"\s*([a-z-]+)", segment, re.I)
            # A shorthand with no property name transitions `all`, per spec.
            prop = head.group(1).lower() if head and not _TIME.match(segment.strip()) else "all"
            if prop in _LAYOUT_TRANSITION_PROPS:
                scan.layout_transitions.append(site)
                break


def _scan_tailwind(text: str, label: str, scan: _SourceScan) -> None:
    def site(m: re.Match) -> _Site:
        return _Site(label, _line_of(text, m.start()), m.group(0)[:48])

    for m in _TW_SPACING.finditer(text):
        scan.spacing_total += 1
        if _off_grid(float(m.group(1))):
            scan.spacing_off.append(site(m))
    for m in _TW_TEXT_PX.finditer(text):
        scan.sizes.setdefault(round(float(m.group(1)), 1), site(m))
    for m in _TW_TEXT_STEP.finditer(text):
        scan.sizes.setdefault(_TW_FONT_SIZE_PX[m.group(1)], site(m))
    for m in _TW_WEIGHT.finditer(text):
        scan.weights.setdefault(_TW_FONT_WEIGHT[m.group(1)], site(m))
    for m in _TW_RADIUS.finditer(text):
        scan.radii.setdefault(f"{m.group(1)}px", site(m))
    for m in _TW_DURATION.finditer(text):
        scan.has_motion = True
        scan.durations += 1
        if float(m.group(1) or m.group(2)) > MOTION_MAX_MS:
            scan.durations_over.append(site(m))
    for m in _TW_TRANSITION_ALL.finditer(text):
        scan.has_motion = True
        scan.layout_transitions.append(site(m))
    for m in _TW_ANIMATE.finditer(text):
        if not _LOADER_HINT.search(m.group(1)):
            scan.infinite_ambient.append(site(m))
    if _TW_MOTION.search(text):
        scan.has_motion = True
    for m in _TW_COLOR.finditer(text):
        scan.colors.append(site(m))
    for m in _TW_OUTLINE_NONE.finditer(text):
        scan.kills_outline.append(site(m))
    if _TW_FOCUS_RESTORE.search(text):
        scan.restores_focus = True


def _scan_keyframes(text: str, label: str, scan: _SourceScan) -> None:
    for m in _KEYFRAMES_OPEN.finditer(text):
        scan.keyframes += 1
        depth, j = 1, m.end()
        while j < len(text) and depth:
            if text[j] == "{":
                depth += 1
            elif text[j] == "}":
                depth -= 1
            j += 1
        block = text[m.end():j]
        for t in _TRANSLATE.finditer(block):
            if abs(float(t.group(1))) > MOTION_TRANSLATE_MAX_PX:
                scan.translate_over.append(_Site(
                    label, _line_of(text, m.end() + t.start()),
                    f"@keyframes {m.group(1)}: {t.group(0)}"))
                break


def _scan_copy(text: str, label: str, suffix: str, scan: _SourceScan) -> None:
    if suffix in SOURCE_MARKUP_SUFFIXES:
        prose = _SRC_TAG.sub(lambda t: re.sub(r"[^\n]", " ", t.group(0)),
                             _SRC_SCRIPT_BLOCK.sub(lambda t: re.sub(r"[^\n]", " ", t.group(0)),
                                                   _STYLE_BLOCK.sub(lambda t: re.sub(r"[^\n]", " ", t.group(0)), text)))
    elif suffix in SOURCE_SCRIPT_SUFFIXES:
        prose = text                      # JSX text and string literals both count
    else:
        return
    scan.markup_files += 1
    for m in _SLOP.finditer(prose):
        scan.slop.append(_Site(label, _line_of(prose, m.start()), f"“{m.group(0)}”"))


def _scan_runtime(text: str, label: str, scan: _SourceScan) -> None:
    """The motion-runtime checks: run on every script-bearing file."""
    scan.script_files += 1
    if _VISIBILITY_GATE.search(text):
        scan.has_visibility_gate = True
    for m in _WHEEL_LISTENER.finditer(text):
        if _PREVENT_DEFAULT.search(text):
            scan.scrolljack.append(_Site(label, _line_of(text, m.start()),
                                         m.group(0)[:40] + " + preventDefault()"))
    for m in _SCROLLJACK_LIB.finditer(text):
        scan.scrolljack.append(_Site(label, _line_of(text, m.start()), m.group(0)))
    for m in _DPR_USE.finditer(text):
        window = text[max(0, m.start() - 80):m.end() + 80]
        if _DPR_CAP.search(window):
            scan.dpr_capped += 1
        else:
            scan.dpr_uncapped.append(_Site(label, _line_of(text, m.start()),
                                           text[m.start():m.end() + 24].split("\n")[0]))
    loops = list(_RENDER_LOOP.finditer(text))
    # One requestAnimationFrame is a "next frame" hop; a loop calls it from
    # inside itself, so it shows up at least twice. setAnimationLoop is a loop
    # by definition.
    if any("setAnimationLoop" in m.group(0) for m in loops) or len(loops) >= 2:
        scan.render_loops.append(_Site(label, _line_of(text, loops[0].start()),
                                       f"{len(loops)} render-loop call(s)"))


def _emit_scale(report: Report, check: str, tag: str, found: dict, cap: int, noun: str,
                key_fmt=str) -> None:
    if not found:
        report.add(check, "SKIP", tag, f"no {noun} declarations found")
        return
    keys = sorted(found, key=lambda k: (isinstance(k, str), k))
    listed = ", ".join(key_fmt(k) for k in keys)
    if len(found) > cap:
        report.add(check, "WARN", tag,
                   f"{len(found)} distinct {noun} values ({listed}) — the rule allows {cap}; "
                   f"first sites: {_first([found[k] for k in keys])}",
                   measured=len(found), threshold=cap)
    else:
        report.add(check, "PASS", tag, f"{len(found)} distinct {noun} values ({listed})",
                   measured=len(found), threshold=cap)


def lint_source(paths: list[Path], root: Path, report: Report) -> None:
    """Scan a UI source tree once; the type scale and motion guard are
    product-level properties, so findings aggregate across files and each
    example carries its own `file:line`."""
    texts: dict[str, str] = {}
    for path in paths:
        try:
            label = str(path.relative_to(root))
        except ValueError:
            label = path.name
        try:
            texts[label] = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            report.add("source.read", "SKIP", label, f"unreadable: {exc}")
    tag = str(root)
    if not texts:
        report.add("source.scope", "SKIP", tag, "no readable UI source files")
        return

    props: dict[str, str] = {}
    for raw in texts.values():
        props.update(custom_properties(raw))
    scan = _SourceScan()
    by_suffix: dict[str, int] = {}
    for label, raw in texts.items():
        suffix = Path(label).suffix.lower()
        by_suffix[suffix] = by_suffix.get(suffix, 0) + 1
        text = _blank_comments(raw)
        if _MOTION_GUARD.search(text):
            scan.has_guard = True
        if suffix in SOURCE_RUNTIME_SUFFIXES:
            _scan_runtime(text, label, scan)       # runtime checks only, by design
            continue
        if _FOCUS_RESTORE.search(text):
            scan.restores_focus = True
        _scan_declarations(text, label, props, scan)
        _scan_tailwind(text, label, scan)
        _scan_keyframes(text, label, scan)
        _scan_copy(text, label, suffix, scan)
        if suffix in SOURCE_SCRIPT_SUFFIXES or suffix in SOURCE_MARKUP_SUFFIXES:
            _scan_runtime(text, label, scan)

    report.add("source.scope", "PASS", tag,
               f"{len(texts)} file(s): " + ", ".join(f"{n} {s}" for s, n in sorted(by_suffix.items())))

    # -- spacing grid ---------------------------------------------------------
    if scan.spacing_total == 0:
        report.add("source.spacing.grid", "SKIP", tag, "no px/rem spacing declarations found")
    elif scan.spacing_off:
        report.add("source.spacing.grid", "WARN", tag,
                   f"{len(scan.spacing_off)} of {scan.spacing_total} spacing values are off the "
                   f"{SPACING_GRID_PX}-px grid (margins {scan.margins}, gaps {scan.gaps}): "
                   f"{_first(scan.spacing_off)} — each needs a written visual reason or the "
                   f"nearest scale step", measured=len(scan.spacing_off), threshold=0)
    else:
        report.add("source.spacing.grid", "PASS", tag,
                   f"{scan.spacing_total} spacing values, all on the {SPACING_GRID_PX}-px grid "
                   f"(margins {scan.margins}, gaps {scan.gaps}; 1–{SPACING_NUDGE_PX} px nudges ignored)",
                   measured=0, threshold=0)

    # -- type scale -----------------------------------------------------------
    _emit_scale(report, "source.type.sizes", tag, scan.sizes, TYPE_SIZES_MAX,
                "font-size", key_fmt=lambda k: f"{k:g}px")
    _emit_scale(report, "source.type.weights", tag, scan.weights, TYPE_WEIGHTS_MAX, "font-weight")
    _emit_scale(report, "source.type.families", tag, scan.families, TYPE_FAMILIES_MAX,
                "font-family")

    # -- geometry -------------------------------------------------------------
    _emit_scale(report, "source.geometry.radius", tag, scan.radii, RADIUS_VALUES_MAX,
                "border-radius")
    _emit_scale(report, "source.geometry.shadow", tag, scan.shadows, SHADOW_VALUES_MAX,
                "box-shadow", key_fmt=lambda k: k[:32])

    # -- motion ---------------------------------------------------------------
    if not scan.has_motion:
        report.add("source.motion", "SKIP", tag, "no transition or animation declared")
    else:
        if scan.durations_over:
            report.add("source.motion.duration", "WARN", tag,
                       f"{len(scan.durations_over)} of {scan.durations} durations exceed "
                       f"{MOTION_MAX_MS:.0f} ms (reserved for one signature sequence): "
                       f"{_first(scan.durations_over)}",
                       measured=len(scan.durations_over), threshold=0)
        elif scan.durations:
            report.add("source.motion.duration", "PASS", tag,
                       f"all {scan.durations} durations at or under {MOTION_MAX_MS:.0f} ms",
                       measured=0, threshold=0)
        if scan.layout_transitions:
            report.add("source.motion.property", "WARN", tag,
                       f"{len(scan.layout_transitions)} transition(s) on a layout property or "
                       f"`all` (relayout on every frame): {_first(scan.layout_transitions)}",
                       measured=len(scan.layout_transitions), threshold=0)
        else:
            report.add("source.motion.property", "PASS", tag,
                       "transitions name transform/opacity-class properties only")
        if scan.infinite_ambient:
            report.add("source.motion.infinite", "WARN", tag,
                       f"{len(scan.infinite_ambient)} infinite animation(s) that are not loaders: "
                       f"{_first(scan.infinite_ambient)}",
                       measured=len(scan.infinite_ambient), threshold=0)
        else:
            report.add("source.motion.infinite", "PASS", tag,
                       "no ambient infinite animation (loaders excepted)")
        if scan.translate_over:
            report.add("source.motion.displacement", "WARN", tag,
                       f"{len(scan.translate_over)} keyframe(s) move more than "
                       f"{MOTION_TRANSLATE_MAX_PX:.0f} px: {_first(scan.translate_over)}",
                       measured=len(scan.translate_over), threshold=MOTION_TRANSLATE_MAX_PX)
        elif scan.keyframes:
            report.add("source.motion.displacement", "PASS", tag,
                       f"{scan.keyframes} @keyframes, none translating past "
                       f"{MOTION_TRANSLATE_MAX_PX:.0f} px")
        if scan.has_guard:
            report.add("source.motion.reduced", "PASS", tag,
                       "motion is guarded by prefers-reduced-motion")
        else:
            report.add("source.motion.reduced", "FAIL", tag,
                       "declares transition/animation with no prefers-reduced-motion guard "
                       "anywhere under this root — if the guard lives in a global stylesheet "
                       "outside it, run on the directory that contains both")

    # -- motion runtime (signature-motion.md §5) ------------------------------
    if not scan.script_files:
        report.add("source.motion.runtime", "SKIP", tag,
                   "no script-bearing files — scroll-jacking, DPR and render-loop checks did not run")
    else:
        if scan.scrolljack:
            report.add("source.motion.scrolljack", "WARN", tag,
                       f"{len(scan.scrolljack)} site(s) intercept the user's scroll: "
                       f"{_first(scan.scrolljack)} — pin the scene, never decide the page's "
                       f"scroll for the reader (a drag handler is the reviewer's call)",
                       measured=len(scan.scrolljack), threshold=0)
        else:
            report.add("source.motion.scrolljack", "PASS", tag,
                       f"no wheel/touch interception or page-snapping library in "
                       f"{scan.script_files} script-bearing file(s)")
        if scan.dpr_uncapped:
            report.add("source.motion.dpr", "WARN", tag,
                       f"{len(scan.dpr_uncapped)} devicePixelRatio use(s) with no cap nearby: "
                       f"{_first(scan.dpr_uncapped)} — cap at 2 (Math.min) or a 3x canvas "
                       f"drops frames on phones", measured=len(scan.dpr_uncapped), threshold=0)
        elif scan.dpr_capped:
            report.add("source.motion.dpr", "PASS", tag,
                       f"devicePixelRatio capped at all {scan.dpr_capped} use(s)")
        else:
            report.add("source.motion.dpr", "SKIP", tag, "devicePixelRatio not used")
        if scan.render_loops and not scan.has_visibility_gate:
            report.add("source.motion.render-loop", "WARN", tag,
                       f"{len(scan.render_loops)} render loop(s) with no visibility gate "
                       f"anywhere under this root: {_first(scan.render_loops)} — pause "
                       f"off-screen (IntersectionObserver, visibilitychange) and stop after "
                       f"the intro", measured=len(scan.render_loops), threshold=0)
        elif scan.render_loops:
            report.add("source.motion.render-loop", "PASS", tag,
                       f"{len(scan.render_loops)} render loop(s) and a visibility gate is present")
        else:
            report.add("source.motion.render-loop", "SKIP", tag, "no render loop found")

    # -- focus visibility -----------------------------------------------------
    if not scan.kills_outline:
        report.add("source.focus", "SKIP", tag, "nothing removes the focus outline")
    elif scan.restores_focus:
        report.add("source.focus", "PASS", tag,
                   f"outline removed at {len(scan.kills_outline)} site(s) and a focus style is restored")
    else:
        report.add("source.focus", "FAIL", tag,
                   f"removes the focus outline without a replacement indicator: "
                   f"{_first(scan.kills_outline)}")

    # -- token discipline -----------------------------------------------------
    if scan.colors:
        report.add("source.color.literal", "WARN", tag,
                   f"{len(scan.colors)} colour literal(s) outside custom-property definitions "
                   f"(a token file is the fix): {_first(scan.colors)}",
                   measured=len(scan.colors), threshold=0)
    else:
        report.add("source.color.literal", "PASS", tag,
                   "no colour literal outside custom-property definitions")

    # -- copy -----------------------------------------------------------------
    if not scan.markup_files:
        report.add("source.copy.slop", "SKIP", tag, "no markup or JSX files to read copy from")
    elif scan.slop:
        report.add("source.copy.slop", "WARN", tag,
                   f"{len(scan.slop)} banned marketing phrase(s): {_first(scan.slop)} — say what "
                   f"happens, with a number", measured=len(scan.slop), threshold=0)
    else:
        report.add("source.copy.slop", "PASS", tag,
                   f"no banned marketing phrase in {scan.markup_files} markup/JSX file(s)")

    # -- what a static scan cannot see ----------------------------------------
    report.add("source.unmeasured", "SKIP", tag,
               "one primary action per viewport, shared edges, tier contrast, control "
               "heights, responsive re-composition and per-state completeness are not "
               "visible statically — run the screenshot review in design-review.md")


# ─── Driver ───


IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


def collect(target: Path, lane: str) -> list[Path]:
    if target.is_file():
        return [target]
    if lane == "render":
        return sorted(p for p in target.rglob("*") if p.suffix.lower() in IMAGE_SUFFIXES)
    if lane == "html":
        return sorted(target.rglob("*.html"))
    if lane == "source":
        wanted = SOURCE_SUFFIXES | SOURCE_RUNTIME_SUFFIXES
        return sorted(
            p for p in target.rglob("*")
            if p.is_file() and p.suffix.lower() in wanted
            and not (SOURCE_SKIP_DIRS & set(p.relative_to(target).parts[:-1]))
            and not p.name.endswith((".min.css", ".min.js", ".d.ts")))
    return sorted(target.rglob("*.pptx"))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Measure rendered design output against the Clade design floors.")
    parser.add_argument("lane", choices=["deck", "render", "html", "source"])
    parser.add_argument("target", type=Path, help="file or directory")
    parser.add_argument("--json", action="store_true", help="emit findings as JSON")
    parser.add_argument("--quiet", action="store_true", help="only show FAIL/WARN")
    args = parser.parse_args(argv)

    if not args.target.exists():
        print(f"design-lint: no such path: {args.target}", file=sys.stderr)
        return 2

    paths = collect(args.target, args.lane)
    if not paths:
        print(f"design-lint: nothing to check under {args.target}", file=sys.stderr)
        return 2

    report = Report(lane=args.lane)
    root = args.target if args.target.is_dir() else args.target.parent
    if args.lane == "source":
        # One scan for the whole tree: the type scale, the motion guard and the
        # focus restore are product-level properties, not per-file ones.
        try:
            lint_source(paths, root, report)
        except Exception as exc:  # a crashed check must be visible, never silent
            report.add("source.error", "FAIL", str(root),
                       f"check crashed: {type(exc).__name__}: {exc}")
    else:
        runner = {"deck": lint_deck, "render": lint_render, "html": lint_html}[args.lane]
        for path in paths:
            # Label by path relative to the scan root. Artifact pages are all
            # named index.html, so a basename label silently merges hundreds of
            # distinct findings into one target and makes every aggregate
            # count wrong.
            try:
                label = str(path.relative_to(root))
            except ValueError:
                label = path.name
            try:
                runner(path, report, label)
            except Exception as exc:  # a crashed check must be visible, never silent
                report.add(f"{args.lane}.error", "FAIL", label,
                           f"check crashed: {type(exc).__name__}: {exc}")

    if args.json:
        print(json.dumps({"lane": report.lane,
                          "failed": report.failed,
                          "findings": [asdict(f) for f in report.findings]}, indent=2))
    else:
        shown = report.findings
        if args.quiet:
            shown = [f for f in shown if f.severity in ("FAIL", "WARN")]
        filtered = Report(lane=report.lane, findings=shown)
        print(f"design-lint {args.lane}: {len(paths)} target(s)")
        print(filtered.render() if shown else "  (no findings)")
        if args.quiet and not shown:
            print(f"  ── {report.lane}: all checks clear")

    return 1 if report.failed else 0


if __name__ == "__main__":
    sys.exit(main())
